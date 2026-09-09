# -*- coding: utf-8 -*-
import openpyxl, html, collections, re
UP="/root/.claude/uploads/12e5ca2d-7231-5774-9600-1c84384c6276/"
E=html.escape

ROOMS=[
 dict(key="试剂准备间", badge="试剂准备间", zone="试剂准备区", sign="试剂准备", z=1,
   what="配制与分装 PCR 反应体系，储存与分装试剂。全流程最洁净的一间，只出不进——试剂从这里发往下游，样本和文库永不进入本间。",
   cond="房间独立、带缓冲间，位于分区最上游，与所有含样本的房间物理隔开。",
   note="纯水仪需就近给排水；超净工作台 150×75 cm 建议独占一面墙。",
   gear=[("超净工作台 ACB-1300V",1),("纯水仪 18.2 MΩ",1),("冷藏冷冻箱 2–8 / −20 ℃",1),
         ("漩涡振荡器 · 掌式离心机",2),("单道移液器 6 支",6)],
   more="另含计时器、温湿度计 2 项"),
 dict(key="PCR2", badge="PCR2", zone="标本与文库制备区", sign="标本制备", z=2,
   what="石蜡切片、组织样本的核酸提取与文库构建。全流程唯一开放操作生物样本的环节，配 A2 型生物安全柜。",
   cond="独立成间、带缓冲，与试剂准备区完全分开，符合标本制备区的隔离要求。",
   note="生物安全柜整机高 227 cm，设备进场前确认吊顶净高与排风接口。",
   gear=[("生物安全柜 A2 型",1),("PCR 扩增仪 ETC-821M",1),("高速离心机 HT165R",1),
         ("恒温混匀仪 96 孔",1),("磁力架 16 孔 · 96 孔",2)],
   more="另含冷藏冷冻箱、漩涡振荡器、移液器等 14 项"),
 dict(key="PCR3", badge="PCR3", zone="血浆文库制备区", sign="文库", z=3,
   what="血浆游离 DNA（cfDNA）的文库构建。之所以要单独一间，是因为组织样本的 DNA 浓度远高于血浆中的 ctDNA，同室操作会把低丰度突变淹掉。",
   cond="与 PCR2 相邻但完全分室，两间设备各自独立配置，不共用器具。",
   note="设备配置与 PCR2 完全一致；两间的移液器、离心机等器具不得互借。",
   gear=[("生物安全柜 A2 型",1),("PCR 扩增仪 ETC-821M",1),("高速离心机 HT165R",1),
         ("恒温混匀仪 96 孔",1),("磁力架 16 孔 · 96 孔",2)],
   more="另含冷藏冷冻箱、漩涡振荡器、移液器等 14 项"),
 dict(key="PCR4", badge="PCR4", zone="杂交捕获区", sign="扩增", z=4,
   what="探针杂交、磁珠捕获与洗脱，把目标基因区域从全基因组文库里富集出来。设备品类最多的一间（21 项），全是台面器械，没有大型落地设备。",
   cond="房间尺寸容得下 21 项台面器械，无需为大型设备预留落地空间。",
   note="器械多，建议沿两面墙 L 形布置台面。<b>现场门牌写的是「扩增」，与本次功能不符，需更换标识。</b>",
   gear=[("PCR 扩增仪 ETC-821M",1),("垂直混匀仪 HS-3",1),("恒温混匀仪 · 高速离心机",2),
         ("荧光定量仪 Fluo-200",1),("磁力架 16 孔 · 96 孔",2)],
   more="另含冷藏冷冻箱、漩涡振荡器、移液器等 14 项"),
 dict(key="PCR5", badge="PCR5", zone="文库扩增与检测区", sign="捕获", z=5,
   what="捕获后文库扩增，以及上机前质检——毛细管电泳看片段分布、荧光定量测浓度，两项都合格才能上机。",
   cond="位于分区下游、独立成间，与上游建库各间物理分离。",
   note="从本间开始有大量扩增产物，人员与物品不得返回上游任一房间。<b>现场门牌写的是「捕获」，需更换标识。</b>",
   gear=[("毛细管电泳仪 Qsep1",1),("荧光定量仪 Fluo-200",1),("PCR 扩增仪 ETC-821M",1),
         ("冷藏冷冻箱 2–8 / −20 ℃",1)],
   more="另含漩涡振荡器、掌式离心机、移液器等 10 项"),
 dict(key="PCR8", badge="PCR8", zone="测序区", sign="文库 QC", z=6,
   what="MGISEQ-2000 上机测序，Halos 一体机完成数据分析与报告解读。全流程终点，设备只有 3 项。",
   cond="位于流程末端，与建库各间物理分离，产物不会反向污染上游。",
   note="对供电稳定与温湿度最敏感——测序中途断电会报废整个 run。清单已配抽湿机控湿，<b>供电保障方式与院内网接入建议一并确认</b>。",
   gear=[("MGISEQ-2000 测序仪",1),("Halos 分析解读一体机",1),("抽湿机 OJ-501E",1)],
   more=""),
]


def _w(t, size=10.0):
    """粗估文本宽度：CJK 全宽，其余半宽"""
    return sum(size if ord(c) > 0x2E80 else size*0.56 for c in t)

def _wrap(t, maxw, size=10.0):
    """按估算宽度折行，最多两行"""
    lines, cur = [], ""
    for ch in t:
        if _w(cur+ch, size) > maxw and cur:
            lines.append(cur); cur = ch
        else:
            cur += ch
    if cur: lines.append(cur)
    return lines[:2]

# ---------- 读取设备清单 ----------
wb=openpyxl.load_workbook(UP+"f37a786e-_____________________.xlsx",data_only=True); ws=wb.active
data=collections.OrderedDict(); room=zone=None
for r in ws.iter_rows(min_row=2,values_only=True):
    a,b,c,d,e,f,g,h=(list(r)+[None]*8)[:8]
    if c=="产品编码": continue
    if a: room,zone=str(a).strip(),(str(b).strip() if b else "")
    if room and d:
        data.setdefault(room,[]).append(dict(code=str(c).strip(),prod=str(d).strip(),
            dev=str(f).strip(),spec=("—" if g in (None,"—","-") else str(g).strip()),qty=int(h or 1)))
for R in ROOMS:
    R["n"]=len(data[R["key"]])
    cov=sum(c for _,c in R["gear"])
    rest=R["n"]-cov
    assert (rest==0)==(R["more"]=="") and (R["more"]=="" or str(rest) in R["more"]), \
        f'{R["key"]}: 未覆盖 {rest} 项，more="{R["more"]}"'
print("设备计数校验通过:", {R["key"]:R["n"] for R in ROOMS})

# ---------- SVG 平面草图 ----------
W,H=1120,462; X0,X1=10,1110; COLW=(X1-X0)/6
CY0,CY1=10,52          # 走廊
BY1=100                # 缓冲下沿
LY1=392                # 实验间下沿
s=[]
s.append(f'<svg viewBox="0 0 {W} {H}" role="img" aria-label="六个功能间沿走廊一字排开的平面草图：自左向右依次为试剂准备区、标本与文库制备区、血浆文库制备区、杂交捕获区、文库扩增与检测区、测序区，每间各有独立缓冲间，工艺流程自左向右单向推进">')
s.append('<defs><marker id="ar" viewBox="0 0 10 10" refX="9" refY="5" markerWidth="7" markerHeight="7" orient="auto-start-reverse">'
         '<path d="M0,0 L10,5 L0,10 z" fill="var(--accent)"/></marker></defs>')
# 底色
s.append(f'<rect x="{X0}" y="{CY0}" width="{X1-X0}" height="{CY1-CY0}" fill="var(--sunk)"/>')
for i,R in enumerate(ROOMS):
    x=X0+i*COLW
    s.append(f'<rect x="{x:.1f}" y="{CY1}" width="{COLW:.1f}" height="{BY1-CY1}" fill="var(--sunk)"/>')
    s.append(f'<rect x="{x:.1f}" y="{BY1}" width="{COLW:.1f}" height="{LY1-BY1}" fill="var(--surface)"/>')
    s.append(f'<rect x="{x:.1f}" y="{BY1}" width="{COLW:.1f}" height="4" fill="var(--z{R["z"]})"/>')
# 墙体（带门洞）
wall='stroke="var(--ink)" stroke-width="2.2" stroke-linecap="square"'
thin='stroke="var(--ink)" stroke-width="1.6"'
s.append(f'<line x1="{X0}" y1="{CY0}" x2="{X1}" y2="{CY0}" {wall}/>')
s.append(f'<line x1="{X0}" y1="{LY1}" x2="{X1}" y2="{LY1}" {wall}/>')
s.append(f'<line x1="{X0}" y1="{CY0}" x2="{X0}" y2="{LY1}" {wall}/>')
s.append(f'<line x1="{X1}" y1="{CY0}" x2="{X1}" y2="{LY1}" {wall}/>')
for yy in (CY1,BY1):
    cur=X0
    for i in range(6):
        x=X0+i*COLW; g0,g1=x+20,x+50
        s.append(f'<line x1="{cur:.1f}" y1="{yy}" x2="{g0:.1f}" y2="{yy}" {thin}/>')
        cur=g1
    s.append(f'<line x1="{cur:.1f}" y1="{yy}" x2="{X1}" y2="{yy}" {thin}/>')
for i in range(1,6):
    x=X0+i*COLW
    s.append(f'<line x1="{x:.1f}" y1="{CY1}" x2="{x:.1f}" y2="{LY1}" {thin}/>')
# 门（合页 + 门扇 + 弧）
for i in range(6):
    x=X0+i*COLW
    for yy in (CY1,BY1):
        s.append(f'<g stroke="var(--line-2)" stroke-width="1.3" fill="none">'
                 f'<line x1="{x+20:.1f}" y1="{yy}" x2="{x+20:.1f}" y2="{yy+30}"/>'
                 f'<path d="M{x+50:.1f},{yy} A30,30 0 0 1 {x+20:.1f},{yy+30}"/></g>')
s.append(f'<text x="{(X0+X1)/2}" y="{CY0+27}" text-anchor="middle" font-size="12.5" letter-spacing="6" '
         f'fill="var(--muted)" font-family="Noto Sans SC,sans-serif">走　廊</text>')
# 房间内容
for i,R in enumerate(ROOMS):
    x=X0+i*COLW; ix=x+12; iw=COLW-24; z=f'var(--z{R["z"]})'
    s.append(f'<text x="{x+COLW/2:.1f}" y="{CY1+31}" text-anchor="middle" font-size="11.5" letter-spacing="2" '
             f'fill="var(--muted)" font-family="Noto Sans SC,sans-serif">缓冲间</text>')
    bw=len(R["badge"])*(10.5 if R["badge"]!="试剂准备间" else 11)+16 if not R["badge"].startswith("PCR") else 46
    s.append(f'<rect x="{ix:.1f}" y="112" width="{bw:.0f}" height="18" rx="2" fill="{z}"/>')
    s.append(f'<text x="{ix+bw/2:.1f}" y="125" text-anchor="middle" font-size="10.5" font-weight="500" '
             f'fill="var(--surface)" font-family="IBM Plex Mono,monospace">{E(R["badge"])}</text>')
    s.append(f'<text x="{x+COLW-12:.1f}" y="125" text-anchor="end" font-size="10.5" fill="var(--muted)" '
             f'font-family="Noto Sans SC,sans-serif">设备 {R["n"]} 项</text>')
    s.append(f'<text x="{ix:.1f}" y="155" font-size="14.5" font-weight="700" fill="var(--ink)" '
             f'font-family="Noto Sans SC,sans-serif">{E(R["zone"])}</text>')
    s.append(f'<text x="{ix:.1f}" y="173" font-size="10.5" fill="var(--faint)" '
             f'font-family="Noto Sans SC,sans-serif">现场门牌：{E(R["sign"])}</text>')
    s.append(f'<line x1="{ix:.1f}" y1="184" x2="{x+COLW-12:.1f}" y2="184" stroke="var(--line)" stroke-width="1"/>')
    s.append(f'<text x="{ix:.1f}" y="201" font-size="10" letter-spacing="1.5" fill="var(--faint)" '
             f'font-family="Noto Sans SC,sans-serif">台面主要设备</text>')
    for j,(lab,_) in enumerate(R["gear"]):
        y=209+j*25
        s.append(f'<rect x="{ix:.1f}" y="{y}" width="{iw:.1f}" height="21" rx="2" fill="var(--z{R["z"]}s)" '
                 f'stroke="{z}" stroke-opacity="0.35" stroke-width="1"/>')
        s.append(f'<text x="{ix+8:.1f}" y="{y+14.5}" font-size="10.5" fill="var(--ink-2)" '
                 f'font-family="Noto Sans SC,sans-serif">{E(lab)}</text>')
    if R["more"]:
        y=209+len(R["gear"])*25+13
        for k,ln in enumerate(_wrap(R["more"], iw-2)):
            s.append(f'<text x="{ix:.1f}" y="{y+k*13}" font-size="10" fill="var(--muted)" '
                     f'font-family="Noto Sans SC,sans-serif">{E(ln)}</text>')
# 流向
s.append(f'<text x="{(X0+X1)/2}" y="424" text-anchor="middle" font-size="11.5" fill="var(--accent)" '
         f'font-family="Noto Sans SC,sans-serif">单向流：标本与文库只向右传递，人员不逆向返回上游</text>')
s.append(f'<line x1="{X0+30}" y1="442" x2="{X1-30}" y2="442" stroke="var(--accent)" stroke-width="2" marker-end="url(#ar)"/>')
s.append('</svg>')
PLAN=('<figure class="plan">'+"".join(s)+
  '<figcaption>依据《设备尺寸参数与摆放示意图》第 2 页重绘，<b>示意图，非按比例平面图</b>。'
  '「现场门牌」为该房间图纸上的原有标注：前三间与本次功能一致，'
  '<b>后三间（杂交捕获 / 文库扩增与检测 / 测序）与现场门牌错位，需更换标识</b>，避免走错房间。</figcaption></figure>')

# ---------- 设备清单表 ----------
rows=[]
for R in ROOMS:
    items=data[R["key"]]; n=sum(i["qty"] for i in items)
    rows.append(f'<tbody><tr class="grp" style="--zc:var(--z{R["z"]}); --zs:var(--z{R["z"]}s)">'
                f'<th colspan="5">{E(R["key"])} · {E(R["zone"])}<span class="mono">{len(items)} 项 / {n} 台件</span></th></tr></tbody>')
    body=[]
    for it in items:
        if it["dev"] in ("单道移液器","八道移液器"): continue
        body.append(f'<tr><td class="code">{E(it["code"])}</td><td>{E(it["prod"])}</td>'
                    f'<td>{E(it["dev"])}</td><td>{E(it["spec"])}</td><td class="num">{it["qty"]}</td></tr>')
    for label in ("单道移液器","八道移液器"):
        grp=[i for i in items if i["dev"]==label]
        if not grp: continue
        rng=[re.sub(r'^.*?移液器','',x["prod"]).strip() for x in grp]
        body.append(f'<tr><td class="code">{E(" / ".join(x["code"] for x in grp))}</td>'
                    f'<td>TIANGEN {E(label)}　<span class="mono" style="font-size:12px;color:var(--muted)">{E(" · ".join(rng))}</span></td>'
                    f'<td>{E(label)}</td><td>—</td><td class="num">{len(grp)}</td></tr>')
    rows.append("<tbody>"+"".join(body)+"</tbody>")

# ---------- 房间卡 ----------
cards=[]
for R in ROOMS:
    cards.append(f'''<article class="card" style="--zc:var(--z{R['z']}); --zs:var(--z{R['z']}s)">
  <div class="chead"><span class="cno">{E(R['badge'])}</span><h3>{E(R['zone'])}</h3>
    <span class="sign"><em>现场门牌</em>　{E(R['sign'])}</span></div>
  <div class="cbody">
    <div><h4>做 什 么</h4><p>{R['what']}</p></div>
    <div class="cond"><span class="pill">现有条件满足</span><p>{R['cond']}</p></div>
    <div class="note"><b style="color:var(--faint);font-weight:500">落位要点　</b>{R['note']}</div>
  </div></article>''')

# ---------- 试剂盒表 ----------
ws2=openpyxl.load_workbook(UP+"165b8495-___________.xlsx",data_only=True).active
kr=[]
for r in ws2.iter_rows(min_row=2,values_only=True):
    no,name,mfr,reg,scope,ext,bar,price=(list(r)+[None]*8)[:8]
    if not no: continue
    kr.append(f'''<tr><td class="num" style="text-align:left">{E(str(no))}</td>
      <td class="w"><b>{E(str(name).strip())}</b><br><span style="font-size:12px;color:var(--muted)">{E(str(mfr).strip())}</span></td>
      <td class="code">{E(str(reg).strip())}</td><td style="min-width:6em">{E(str(scope).strip())}</td>
      <td class="w" style="font-size:12.5px;color:var(--ink-2)">{E((str(ext).strip() if ext else "") or "—")}</td>
      <td class="num"><b>{int(price):,}</b> 元</td></tr>''')
kits=('<thead><tr><th>#</th><th>试剂盒 / 生产厂家</th><th>注册证号</th><th>注册适用范围</th>'
      '<th>可拓展适用范围</th><th style="text-align:right">物价收费</th></tr></thead><tbody>'+"".join(kr)+'</tbody>')

t=open("tpl2.html",encoding="utf-8").read()
for k,v in (("__PLAN__",PLAN),("__CARDS__","\n".join(cards)),("__EQUIP__","\n".join(rows)),("__KITS__",kits)):
    assert k in t, k
    t=t.replace(k,v)
open("ngs-lab.html","w",encoding="utf-8").write(t)
print("written", len(t), "bytes")
