# -*- coding: utf-8 -*-
"""打印稿：黑白、A4 竖版、三线表。数据源与投屏版同一份 xlsx。"""
import openpyxl, html, collections, re
UP="/root/.claude/uploads/12e5ca2d-7231-5774-9600-1c84384c6276/"
E=html.escape

ROOMS=[
 dict(key="试剂准备间", zone="试剂准备区", room="试剂准备间",
  work="配制与分装 PCR 反应体系，储存与分装试剂。全流程最洁净的一间，只出不进，样本与文库不得进入。",
  cond="独立成间并设缓冲间，位于分区最上游，与所有含样本的房间物理隔开。满足要求。"),
 dict(key="PCR2", zone="标本与文库制备区", room="PCR2",
  work="石蜡切片、组织样本的核酸提取与文库构建，配 A2 型生物安全柜，为全流程唯一开放操作生物样本的环节。",
  cond="独立成间并设缓冲间，与试剂准备区完全分开，符合标本制备区的隔离要求。满足要求。"),
 dict(key="PCR3", zone="血浆文库制备区", room="PCR3",
  work="血浆游离 DNA（cfDNA）的文库构建。单独设间是因为组织样本的 DNA 浓度远高于血浆中的 ctDNA，同室操作会淹没低丰度突变。",
  cond="与 PCR2 相邻但完全分室，两间设备各自独立配置、不共用器具。满足要求。"),
 dict(key="PCR4", zone="杂交捕获区", room="PCR4",
  work="探针杂交、磁珠捕获与洗脱，将目标基因区域从全基因组文库中富集出来。",
  cond="全部为台面器械，无大型落地设备，房间尺寸满足器械摆放需要。满足要求。"),
 dict(key="PCR5", zone="文库扩增与检测区", room="PCR5",
  work="捕获后文库扩增，以及上机前质检：毛细管电泳检测片段分布、荧光定量检测浓度，两项合格方可上机。",
  cond="位于分区下游、独立成间，与上游建库各间物理分离。自本间起有大量扩增产物，人员与物品不得返回上游。满足要求。"),
 dict(key="PCR8", zone="测序区", room="PCR8",
  work="上机测序，并由分析解读一体机完成数据分析与报告解读，为全流程终点。",
  cond="位于流程末端，与建库各间物理分离，扩增产物不会反向污染上游。满足要求。"),
]

REF=[
 ("国械注准20193400099","人类BRCA1基因和BRCA2基因突变检测试剂盒（可逆末端终止测序法）","厦门艾德生物医药科技股份有限公司",
  "卵巢癌、乳腺癌",
  "检测 BRCA1 / BRCA2 编码区、外显子-内含子连接区、UTR 区与启动子区的点突变、插入缺失及纯合缺失；用于帕米帕利等 PARP 抑制剂的用药指导"),
 ("国械注准20253402685","人BRCA1/BRCA2基因突变检测试剂盒（可逆末端终止测序法）","厦门艾德生物医药科技股份有限公司",
  "前列腺癌",
  "用于携带胚系和／或体系 BRCA 基因突变的转移性去势抵抗性前列腺癌（mCRPC）；泽倍珂®（尼拉帕利阿比特龙片）的伴随诊断"),
]

DIRS=[
 ("实体瘤精准诊疗","靶向与免疫治疗的伴随诊断：驱动基因突变、基因融合、肿瘤突变负荷（TMB）、微卫星不稳定（MSI）与同源重组缺陷（HRD）。国内已注册产品最多的方向，本次项目即属此类。","已有注册产品"),
 ("遗传性肿瘤易感基因","BRCA1／BRCA2 等胚系突变筛查，用于高危人群管理、亲属级联筛查与 PARP 抑制剂用药指导。表 4 所列两项即属此类。","已有注册产品"),
 ("血液系统肿瘤分子分型","白血病、淋巴瘤、骨髓增生异常综合征的融合基因与突变谱检测，用于分型、危险度分层与疗效监测。微小残留病（MRD）监测目前多以自建项目开展。","部分已有注册产品"),
 ("生殖健康与出生缺陷防控","无创产前筛查（NIPT）、染色体拷贝数变异检测（CNV-seq）、胚胎植入前遗传学检测（PGT）。NIPT 是国内最早获批的 NGS 临床应用方向。","已有注册产品"),
 ("单基因遗传病与携带者筛查","地中海贫血、遗传性耳聋等常见单基因病已有注册试剂盒；更广的多基因 panel 与全外显子组测序目前多以自建项目路径开展。","部分已有注册产品"),
 ("感染性疾病病原检测","宏基因组测序（mNGS）与靶向测序（tNGS），用于疑难危重感染的病原快速识别。测序平台已取得三类注册证，国内已发布多部临床应用专家共识。","多以 LDT 路径开展"),
]

# ---------- 读取设备清单 ----------
ws=openpyxl.load_workbook(UP+"f37a786e-_____________________.xlsx",data_only=True).active
data=collections.OrderedDict(); room=None
for r in ws.iter_rows(min_row=2,values_only=True):
    a,b,c,d,e,f,g,h=(list(r)+[None]*8)[:8]
    if c=="产品编码": continue
    if a: room=str(a).strip()
    if room and d:
        data.setdefault(room,[]).append(dict(code=str(c).strip(),prod=str(d).strip(),
            dev=str(f).strip(),spec=("—" if g in (None,"—","-") else str(g).strip()),qty=int(h or 1)))
for R in ROOMS: R["n"]=len(data[R["key"]])
total=sum(R["n"] for R in ROOMS)
assert total==91, total
print("设备合计校验:", total)

# ---------- 表 1 ----------
rooms_html="".join(
  f'<tr><td class="c">{i}</td><td>{E(R["zone"])}<br>（{E(R["room"])}）</td>'
  f'<td>{E(R["work"])}</td><td>{E(R["cond"])}</td><td class="c">{R["n"]} 项</td></tr>'
  for i,R in enumerate(ROOMS,1))

# ---------- 表 2 ----------
eq=[]
for R in ROOMS:
    items=data[R["key"]]
    eq.append(f'<tr class="grp"><td colspan="5">{E(R["room"])}　{E(R["zone"])}　（{len(items)} 项 / '
              f'{sum(i["qty"] for i in items)} 台件）</td></tr>')
    for it in items:
        if it["dev"] in ("单道移液器","八道移液器"): continue
        eq.append(f'<tr><td class="nw">{E(it["code"])}</td><td>{E(it["prod"])}</td>'
                  f'<td>{E(it["dev"])}</td><td>{E(it["spec"])}</td><td class="c">{it["qty"]}</td></tr>')
    for label in ("单道移液器","八道移液器"):
        grp=[i for i in items if i["dev"]==label]
        if not grp: continue
        rng=[re.sub(r'^.*?移液器','',x["prod"]).strip() for x in grp]
        eq.append(f'<tr><td>{E(" / ".join(x["code"] for x in grp))}</td>'
                  f'<td>TIANGEN {E(label)}（{E("；".join(rng))}）</td>'
                  f'<td>{E(label)}</td><td>—</td><td class="c">{len(grp)}</td></tr>')

# ---------- 表 3 ----------
ws2=openpyxl.load_workbook(UP+"165b8495-___________.xlsx",data_only=True).active
kits=[]
for r in ws2.iter_rows(min_row=2,values_only=True):
    no,name,mfr,reg,scope,ext,bar,price=(list(r)+[None]*8)[:8]
    if not no: continue
    kits.append(f'<tr><td class="c">{E(str(no))}</td>'
                f'<td>{E(str(name).strip())}<br>{E(str(mfr).strip())}</td>'
                f'<td class="nw">{E(str(reg).strip())}</td><td>{E(str(scope).strip())}</td>'
                f'<td>{E((str(ext).strip() if ext else "") or "—")}</td>'
                f'<td class="n">{int(price):,} 元</td></tr>')

# ---------- 表 4 / 表 5 ----------
ref="".join(f'<tr><td class="nw">{E(a)}</td><td>{E(b)}<br>{E(c)}</td><td>{E(d)}</td><td>{E(e)}</td></tr>'
            for a,b,c,d,e in REF)
dirs="".join(f'<tr><td class="c">{i}</td><td>{E(t)}</td><td>{E(d)}</td><td>{E(st)}</td></tr>'
             for i,(t,d,st) in enumerate(DIRS,1))

t=open("tpl_print.html",encoding="utf-8").read()
for k,v in (("__ROOMS__",rooms_html),("__EQUIP__","".join(eq)),("__KITS__","".join(kits)),
            ("__REF__",ref),("__DIRS__",dirs)):
    assert k in t, k
    t=t.replace(k,v)

# 守卫：变量必须定义；打印稿里不得出现彩色声明
_root=re.search(r":root\{(.*?)\}", t, re.S).group(1)
_def=set(re.findall(r"(--[a-z0-9-]+)\s*:", _root))
_used=set(re.findall(r"var\((--[a-z0-9-]+)\)", t))
assert not (_used-_def), f"CSS 变量未定义：{sorted(_used-_def)}"
_css=t[t.index("<style>"):t.index("</style>")]
_hex=set(re.findall(r"#([0-9A-Fa-f]{6})\b", _css))
_colored=[h for h in _hex if not (h[0:2].lower()==h[2:4].lower()==h[4:6].lower())]
assert not _colored, f"打印稿出现彩色：{_colored}"
print("守卫通过：无彩色（全部为中性灰阶）")

open("ngs-report-print.html","w",encoding="utf-8").write(t)
print("written", len(t), "bytes")
